"""
Excel Reporter - Xuat bao cao BCTC ra file Excel (.xlsx).

Tao workbook nhieu sheet:
  1. Thong tin DN
  2. CDKT (Bang Can doi Ke toan)
  3. KQHDKD (Ket qua Hoat dong Kinh doanh)
  4. Doi chieu cheo (Cross-check results)
  5. Bat thuong (Anomalies)
  6. Chi so tai chinh (Ratios)
"""

import io
from typing import Dict, Any, Optional

from openpyxl import Workbook
from openpyxl.styles import (
    Font, Alignment, PatternFill, Border, Side, numbers
)
from openpyxl.utils import get_column_letter


# ─── Style constants ────────────────────────────────────

_HEADER_FONT = Font(name="Arial", bold=True, size=12, color="FFFFFF")
_HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
_SUBHEADER_FONT = Font(name="Arial", bold=True, size=10)
_SUBHEADER_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
_NORMAL_FONT = Font(name="Arial", size=10)
_BOLD_FONT = Font(name="Arial", bold=True, size=10)
_TITLE_FONT = Font(name="Arial", bold=True, size=14, color="2F5496")

_THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

_PASS_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
_FAIL_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
_WARN_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")

_VND_FORMAT = '#,##0'
_PCT_FORMAT = '0.00%'


class ExcelReporter:
    """Tao bao cao Excel tu ket qua phan tich BCTC."""

    def __init__(self, analysis_data: Dict[str, Any]):
        """
        Args:
            analysis_data: Du lieu phan tich day du (response tu /api/upload).
        """
        self.data = analysis_data
        self.company = analysis_data.get("company", {})
        self.reports = analysis_data.get("reports", {})
        self.validation = analysis_data.get("validation", {})
        self.anomalies = analysis_data.get("anomalies", [])
        self.ratios = analysis_data.get("ratios", [])
        self.journals = analysis_data.get("journals", [])
        self.wb = Workbook()

    def generate(self) -> io.BytesIO:
        """Tao workbook va tra ve BytesIO stream."""
        # Xoa sheet mac dinh
        default_sheet = self.wb.active
        self.wb.remove(default_sheet)

        self._create_company_sheet()
        self._create_cdkt_sheet()
        self._create_kqhdkd_sheet()
        self._create_crosscheck_sheet()
        self._create_anomaly_sheet()
        self._create_ratio_sheet()

        output = io.BytesIO()
        self.wb.save(output)
        output.seek(0)
        return output

    # ─── Sheet 1: Thong tin DN ───────────────────────────

    def _create_company_sheet(self):
        ws = self.wb.create_sheet("Thong tin DN")
        ws.sheet_properties.tabColor = "2F5496"

        # Tieu de
        ws.merge_cells("A1:D1")
        cell = ws["A1"]
        cell.value = "THONG TIN DOANH NGHIEP"
        cell.font = _TITLE_FONT
        cell.alignment = Alignment(horizontal="center")

        # Noi dung
        info_rows = [
            ("Ma so thue (MST)", self.company.get("mst", "")),
            ("Ten doanh nghiep", self.company.get("ten_dn", "")),
            ("Dia chi", self.company.get("dia_chi", "")),
            ("Tinh/Thanh pho", self.company.get("tinh", "")),
            ("Ma to khai", self.company.get("ma_to_khai", "")),
            ("Ten to khai", self.company.get("ten_to_khai", "")),
            ("Ky bao cao", self.company.get("ky_bao_cao", "")),
            ("Tu ngay", self.company.get("tu_ngay", "")),
            ("Den ngay", self.company.get("den_ngay", "")),
            ("Ngay lap", self.company.get("ngay_lap", "")),
            ("Thong tu", self.data.get("circular", "")),
        ]

        for idx, (label, value) in enumerate(info_rows, start=3):
            ws.cell(row=idx, column=1, value=label).font = _BOLD_FONT
            ws.cell(row=idx, column=1).border = _THIN_BORDER
            ws.cell(row=idx, column=2, value=value).font = _NORMAL_FONT
            ws.cell(row=idx, column=2).border = _THIN_BORDER

        # Health score
        total_checks = self.validation.get("total_checks", 0)
        passed = self.validation.get("passed", 0)
        health = round((passed / total_checks * 100)) if total_checks > 0 else 0

        row = len(info_rows) + 4
        ws.merge_cells(f"A{row}:D{row}")
        cell = ws.cell(row=row, column=1, value=f"Health Score: {health}%")
        cell.font = Font(name="Arial", bold=True, size=14,
                         color="00B050" if health > 80 else "FFC000" if health > 50 else "FF0000")

        ws.column_dimensions["A"].width = 25
        ws.column_dimensions["B"].width = 50

    # ─── Sheet 2: CDKT ───────────────────────────────────

    def _create_cdkt_sheet(self):
        ws = self.wb.create_sheet("CDKT")
        ws.sheet_properties.tabColor = "00B050"

        cdkt = self.reports.get("cdkt", {})
        cuoi_nam = cdkt.get("so_cuoi_nam", {})
        dau_nam = cdkt.get("so_dau_nam", {})

        # Tieu de
        ws.merge_cells("A1:D1")
        cell = ws["A1"]
        cell.value = "BANG CAN DOI KE TOAN"
        cell.font = _TITLE_FONT
        cell.alignment = Alignment(horizontal="center")

        # Header
        headers = ["Ma chi tieu", "Ten chi tieu", "So cuoi nam (VND)", "So dau nam (VND)"]
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=3, column=col, value=header)
            cell.font = _HEADER_FONT
            cell.fill = _HEADER_FILL
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            cell.border = _THIN_BORDER

        # Import mapper for display names
        try:
            from parser.tt133_mapper import CDKT_MAP
        except ImportError:
            CDKT_MAP = {}

        # Data rows
        row = 4
        sorted_keys = sorted(cuoi_nam.keys())
        for key in sorted_keys:
            val_cn = cuoi_nam.get(key)
            val_dn = dau_nam.get(key)
            info = CDKT_MAP.get(key, {})

            is_bold = info.get("bold", False) or info.get("level", 2) == 0

            ws.cell(row=row, column=1, value=key).font = _BOLD_FONT if is_bold else _NORMAL_FONT
            ws.cell(row=row, column=1).border = _THIN_BORDER

            ws.cell(row=row, column=2, value=info.get("name", key)).font = (
                _BOLD_FONT if is_bold else _NORMAL_FONT
            )
            ws.cell(row=row, column=2).border = _THIN_BORDER

            cell_cn = ws.cell(row=row, column=3)
            if isinstance(val_cn, (int, float)):
                cell_cn.value = val_cn
                cell_cn.number_format = _VND_FORMAT
            else:
                cell_cn.value = val_cn
            cell_cn.font = _BOLD_FONT if is_bold else _NORMAL_FONT
            cell_cn.alignment = Alignment(horizontal="right")
            cell_cn.border = _THIN_BORDER

            cell_dn = ws.cell(row=row, column=4)
            if isinstance(val_dn, (int, float)):
                cell_dn.value = val_dn
                cell_dn.number_format = _VND_FORMAT
            else:
                cell_dn.value = val_dn
            cell_dn.font = _BOLD_FONT if is_bold else _NORMAL_FONT
            cell_dn.alignment = Alignment(horizontal="right")
            cell_dn.border = _THIN_BORDER

            row += 1

        ws.column_dimensions["A"].width = 15
        ws.column_dimensions["B"].width = 45
        ws.column_dimensions["C"].width = 25
        ws.column_dimensions["D"].width = 25

    # ─── Sheet 3: KQHDKD ────────────────────────────────

    def _create_kqhdkd_sheet(self):
        ws = self.wb.create_sheet("KQHDKD")
        ws.sheet_properties.tabColor = "FFC000"

        kqhdkd = self.reports.get("kqhdkd", {})
        nam_nay = kqhdkd.get("nam_nay", {})
        nam_truoc = kqhdkd.get("nam_truoc", {})

        # Tieu de
        ws.merge_cells("A1:D1")
        cell = ws["A1"]
        cell.value = "KET QUA HOAT DONG KINH DOANH"
        cell.font = _TITLE_FONT
        cell.alignment = Alignment(horizontal="center")

        # Header
        headers = ["Ma chi tieu", "Ten chi tieu", "Nam nay (VND)", "Nam truoc (VND)"]
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=3, column=col, value=header)
            cell.font = _HEADER_FONT
            cell.fill = _HEADER_FILL
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            cell.border = _THIN_BORDER

        try:
            from parser.tt133_mapper import KQHDKD_MAP
        except ImportError:
            KQHDKD_MAP = {}

        row = 4
        sorted_keys = sorted(nam_nay.keys())
        for key in sorted_keys:
            val_nn = nam_nay.get(key)
            val_nt = nam_truoc.get(key)
            info = KQHDKD_MAP.get(key, {})

            ws.cell(row=row, column=1, value=key).font = _NORMAL_FONT
            ws.cell(row=row, column=1).border = _THIN_BORDER

            ws.cell(row=row, column=2, value=info.get("name", key)).font = _NORMAL_FONT
            ws.cell(row=row, column=2).border = _THIN_BORDER

            cell_nn = ws.cell(row=row, column=3)
            if isinstance(val_nn, (int, float)):
                cell_nn.value = val_nn
                cell_nn.number_format = _VND_FORMAT
            else:
                cell_nn.value = val_nn
            cell_nn.alignment = Alignment(horizontal="right")
            cell_nn.border = _THIN_BORDER

            cell_nt = ws.cell(row=row, column=4)
            if isinstance(val_nt, (int, float)):
                cell_nt.value = val_nt
                cell_nt.number_format = _VND_FORMAT
            else:
                cell_nt.value = val_nt
            cell_nt.alignment = Alignment(horizontal="right")
            cell_nt.border = _THIN_BORDER

            row += 1

        ws.column_dimensions["A"].width = 15
        ws.column_dimensions["B"].width = 45
        ws.column_dimensions["C"].width = 25
        ws.column_dimensions["D"].width = 25

    # ─── Sheet 4: Doi chieu cheo ─────────────────────────

    def _create_crosscheck_sheet(self):
        ws = self.wb.create_sheet("Doi chieu cheo")
        ws.sheet_properties.tabColor = "FF0000"

        # Tieu de
        ws.merge_cells("A1:E1")
        cell = ws["A1"]
        cell.value = "KET QUA DOI CHIEU CHEO"
        cell.font = _TITLE_FONT
        cell.alignment = Alignment(horizontal="center")

        # Summary
        total = self.validation.get("total_checks", 0)
        passed = self.validation.get("passed", 0)
        warnings = self.validation.get("warnings", 0)
        errors = self.validation.get("errors", 0)

        ws.cell(row=2, column=1, value=f"Tong: {total}").font = _BOLD_FONT
        ws.cell(row=2, column=2, value=f"Dat: {passed}").font = Font(
            name="Arial", bold=True, color="00B050"
        )
        ws.cell(row=2, column=3, value=f"Canh bao: {warnings}").font = Font(
            name="Arial", bold=True, color="FFC000"
        )
        ws.cell(row=2, column=4, value=f"Loi: {errors}").font = Font(
            name="Arial", bold=True, color="FF0000"
        )

        # Header
        headers = ["Ma quy tac", "Ten quy tac", "Ket qua", "Mo ta", "Chenh lech"]
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=4, column=col, value=header)
            cell.font = _HEADER_FONT
            cell.fill = _HEADER_FILL
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            cell.border = _THIN_BORDER

        # Data
        details = self.validation.get("details", [])
        for idx, detail in enumerate(details, start=5):
            severity = detail.get("severity", "")

            ws.cell(row=idx, column=1, value=detail.get("rule_id", "")).font = _NORMAL_FONT
            ws.cell(row=idx, column=1).border = _THIN_BORDER

            ws.cell(row=idx, column=2, value=detail.get("rule_name", "")).font = _NORMAL_FONT
            ws.cell(row=idx, column=2).border = _THIN_BORDER

            result_cell = ws.cell(row=idx, column=3)
            if severity == "ok":
                result_cell.value = "DAT"
                result_cell.fill = _PASS_FILL
            elif severity == "warning":
                result_cell.value = "CANH BAO"
                result_cell.fill = _WARN_FILL
            else:
                result_cell.value = "LOI"
                result_cell.fill = _FAIL_FILL
            result_cell.font = _BOLD_FONT
            result_cell.alignment = Alignment(horizontal="center")
            result_cell.border = _THIN_BORDER

            ws.cell(row=idx, column=4, value=detail.get("message", "")).font = _NORMAL_FONT
            ws.cell(row=idx, column=4).border = _THIN_BORDER

            diff_cell = ws.cell(row=idx, column=5)
            diff = detail.get("difference", 0)
            if isinstance(diff, (int, float)):
                diff_cell.value = diff
                diff_cell.number_format = _VND_FORMAT
            else:
                diff_cell.value = diff
            diff_cell.alignment = Alignment(horizontal="right")
            diff_cell.border = _THIN_BORDER

        ws.column_dimensions["A"].width = 15
        ws.column_dimensions["B"].width = 40
        ws.column_dimensions["C"].width = 15
        ws.column_dimensions["D"].width = 55
        ws.column_dimensions["E"].width = 20

    # ─── Sheet 5: Bat thuong ─────────────────────────────

    def _create_anomaly_sheet(self):
        ws = self.wb.create_sheet("Bat thuong")
        ws.sheet_properties.tabColor = "ED7D31"

        # Tieu de
        ws.merge_cells("A1:F1")
        cell = ws["A1"]
        cell.value = "PHAT HIEN BAT THUONG"
        cell.font = _TITLE_FONT
        cell.alignment = Alignment(horizontal="center")

        ws.cell(row=2, column=1, value=f"Tong so: {len(self.anomalies)} bat thuong").font = _BOLD_FONT

        # Header
        headers = ["Ma", "Muc do", "Danh muc", "Tieu de", "Mo ta", "De xuat"]
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=4, column=col, value=header)
            cell.font = _HEADER_FONT
            cell.fill = _HEADER_FILL
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            cell.border = _THIN_BORDER

        risk_fills = {
            "critical": PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid"),
            "high": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
            "medium": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
            "low": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
        }

        for idx, anomaly in enumerate(self.anomalies, start=5):
            risk_level = anomaly.get("risk_level", "medium")

            ws.cell(row=idx, column=1, value=anomaly.get("code", "")).font = _NORMAL_FONT
            ws.cell(row=idx, column=1).border = _THIN_BORDER

            risk_cell = ws.cell(row=idx, column=2, value=risk_level.upper())
            risk_cell.font = _BOLD_FONT
            risk_cell.fill = risk_fills.get(risk_level, _WARN_FILL)
            risk_cell.alignment = Alignment(horizontal="center")
            risk_cell.border = _THIN_BORDER

            ws.cell(row=idx, column=3, value=anomaly.get("category", "")).font = _NORMAL_FONT
            ws.cell(row=idx, column=3).border = _THIN_BORDER

            ws.cell(row=idx, column=4, value=anomaly.get("title", "")).font = _BOLD_FONT
            ws.cell(row=idx, column=4).border = _THIN_BORDER

            desc_cell = ws.cell(row=idx, column=5, value=anomaly.get("description", ""))
            desc_cell.font = _NORMAL_FONT
            desc_cell.alignment = Alignment(wrap_text=True)
            desc_cell.border = _THIN_BORDER

            sugg_cell = ws.cell(row=idx, column=6, value=anomaly.get("suggestion", ""))
            sugg_cell.font = _NORMAL_FONT
            sugg_cell.alignment = Alignment(wrap_text=True)
            sugg_cell.border = _THIN_BORDER

        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 12
        ws.column_dimensions["C"].width = 12
        ws.column_dimensions["D"].width = 35
        ws.column_dimensions["E"].width = 55
        ws.column_dimensions["F"].width = 45

    # ─── Sheet 6: Chi so tai chinh ───────────────────────

    def _create_ratio_sheet(self):
        ws = self.wb.create_sheet("Chi so tai chinh")
        ws.sheet_properties.tabColor = "7030A0"

        # Tieu de
        ws.merge_cells("A1:G1")
        cell = ws["A1"]
        cell.value = "CHI SO TAI CHINH"
        cell.font = _TITLE_FONT
        cell.alignment = Alignment(horizontal="center")

        # Header
        headers = [
            "Ma", "Ten chi so", "Ten (EN)",
            "Gia tri", "Don vi", "Nguong tham chieu", "Nhan xet"
        ]
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=3, column=col, value=header)
            cell.font = _HEADER_FONT
            cell.fill = _HEADER_FILL
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            cell.border = _THIN_BORDER

        # Group by category
        categories = {
            "liquidity": "Kha nang thanh toan",
            "profitability": "Kha nang sinh loi",
            "leverage": "Co cau von",
            "efficiency": "Hieu qua hoat dong",
        }

        row = 4
        current_category = None

        for ratio in self.ratios:
            cat = ratio.get("category", "")
            if cat != current_category:
                current_category = cat
                ws.merge_cells(f"A{row}:G{row}")
                cat_cell = ws.cell(
                    row=row, column=1,
                    value=categories.get(cat, cat).upper()
                )
                cat_cell.font = _SUBHEADER_FONT
                cat_cell.fill = _SUBHEADER_FILL
                cat_cell.border = _THIN_BORDER
                row += 1

            ws.cell(row=row, column=1, value=ratio.get("code", "")).font = _NORMAL_FONT
            ws.cell(row=row, column=1).border = _THIN_BORDER

            ws.cell(row=row, column=2, value=ratio.get("name", "")).font = _NORMAL_FONT
            ws.cell(row=row, column=2).border = _THIN_BORDER

            ws.cell(row=row, column=3, value=ratio.get("name_en", "")).font = _NORMAL_FONT
            ws.cell(row=row, column=3).border = _THIN_BORDER

            val_cell = ws.cell(row=row, column=4)
            val = ratio.get("value")
            if val is not None:
                val_cell.value = round(val, 4)
                val_cell.number_format = '0.00'
            else:
                val_cell.value = "N/A"
            val_cell.alignment = Alignment(horizontal="right")
            val_cell.font = _BOLD_FONT
            val_cell.border = _THIN_BORDER

            ws.cell(row=row, column=5, value=ratio.get("unit", "")).font = _NORMAL_FONT
            ws.cell(row=row, column=5).alignment = Alignment(horizontal="center")
            ws.cell(row=row, column=5).border = _THIN_BORDER

            ws.cell(row=row, column=6, value=ratio.get("benchmark", "")).font = _NORMAL_FONT
            ws.cell(row=row, column=6).border = _THIN_BORDER

            interp_cell = ws.cell(row=row, column=7, value=ratio.get("interpretation", ""))
            interp_cell.font = _NORMAL_FONT
            interp_cell.alignment = Alignment(wrap_text=True)
            interp_cell.border = _THIN_BORDER

            row += 1

        ws.column_dimensions["A"].width = 10
        ws.column_dimensions["B"].width = 35
        ws.column_dimensions["C"].width = 20
        ws.column_dimensions["D"].width = 12
        ws.column_dimensions["E"].width = 10
        ws.column_dimensions["F"].width = 30
        ws.column_dimensions["G"].width = 30
